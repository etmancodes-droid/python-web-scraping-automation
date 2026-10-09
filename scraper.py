from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side # type: ignore
from openpyxl.worksheet.table import Table, TableStyleInfo # type: ignore


def save_excel(df: pd.DataFrame, file_name: Path) -> None:
    with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Quotes")
        worksheet = writer.sheets["Quotes"]

        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        worksheet.column_dimensions["A"].width = 80
        worksheet.column_dimensions["B"].width = 25
        worksheet.column_dimensions["C"].width = 35

        header_fill = PatternFill("solid", fgColor="17365D")
        header_font = Font(color="FFFFFF", bold=True, size=12)
        stripe_fill = PatternFill("solid", fgColor="EAF2F8")
        border_side = Side(style="thin", color="D6E2F0")

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=Side(style="medium", color="5B9BD5"))

        for row_number in range(2, worksheet.max_row + 1):
            quote_cell = worksheet.cell(row=row_number, column=1)
            tags_cell = worksheet.cell(row=row_number, column=3)
            line_count = max(1, (len(str(quote_cell.value)) + 79) // 80)
            worksheet.row_dimensions[row_number].height = min(18 * line_count, 90)

            for cell in worksheet[row_number]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=border_side)
                if row_number % 2 == 0:
                    cell.fill = stripe_fill

            tags_cell.font = Font(color="2F75B5")

        worksheet.row_dimensions[1].height = 26
        table = Table(
            displayName="QuotesTable",
            ref=f"A1:C{worksheet.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)


def scrape_quotes():
    url = "https://quotes.toscrape.com/"
    file_name = Path(__file__).resolve().with_name("extracted_data.xlsx")
    quotes_data = []

    print("Connecting to the website and scraping all pages...")

    try:
        with requests.Session() as session:
            while url:
                response = session.get(url, timeout=10)
                response.raise_for_status()
                soup = BeautifulSoup(response.text, "html.parser")

                for item in soup.select("div.quote"):
                    text = item.select_one("span.text")
                    author = item.select_one("small.author")
                    if text is None or author is None:
                        raise ValueError("Could not extract a quote or its author.")

                    quotes_data.append({
                        "Quote": text.get_text(strip=True),
                        "Author": author.get_text(strip=True),
                        "Tags": ", ".join(
                            tag.get_text(strip=True) for tag in item.select("a.tag")
                        ),
                    })

                next_page = soup.select_one("li.next a")
                url = urljoin(response.url, next_page["href"]) if next_page else ""

        if not quotes_data:
            raise ValueError("No quotes were found to save.")

        df = pd.DataFrame(quotes_data)
        try:
            save_excel(df, file_name)
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_name = file_name.with_name(f"extracted_data_{timestamp}.xlsx")
            save_excel(df, file_name)

        print(f"Saved {len(quotes_data)} quotes to Excel: {file_name}")

    except requests.exceptions.RequestException as e:
        print(f"Website connection failed: {e}")
    except (OSError, ValueError) as e:
        print(f"Scraping or Excel export failed: {e}")


if __name__ == "__main__":
    scrape_quotes()
